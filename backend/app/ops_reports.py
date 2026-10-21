# app/ops_reports.py
"""Operations portal, sections 15-16: reports with export, and analytics.

Reports share one shape -- {columns, rows} -- so a single table component and
a single exporter serve all of them. Adding a report means adding one builder
function, not another endpoint plus another exporter.
"""
import csv
import io
import logging
from datetime import date, datetime, timedelta
from typing import Any, Callable, Optional

from fastapi import APIRouter, Depends, HTTPException, Query, Request, status
from fastapi.responses import StreamingResponse
from sqlalchemy import func
from sqlalchemy.orm import Session

from . import auth_models as am
from . import ops_models as om
from . import permissions as perm
from .database import get_db

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/ops", tags=["operations"])

MAX_EXPORT_ROWS = 20_000


# ---------------------------------------------------------------------------
# report builders -- each returns (columns, rows)
# ---------------------------------------------------------------------------

def _students(db: Session, start: date, end: date):
    q = (
        db.query(am.User, am.StudentProfile, am.Membership)
        .outerjoin(am.StudentProfile, am.StudentProfile.user_id == am.User.id)
        .outerjoin(am.Membership, am.Membership.user_id == am.User.id)
        .filter(am.User.role == am.ROLE_STUDENT)
        .filter(func.date(am.User.created_at) >= start)
        .filter(func.date(am.User.created_at) <= end)
        .order_by(am.User.created_at.desc())
    )
    cols = ["Member ID", "Name", "Email", "Phone", "College", "Course",
            "Year", "Status", "Registered On"]
    rows = [
        [m.member_id if m else "", p.full_name if p else "", u.email,
         p.phone if p else "", p.college_name if p else "",
         p.course if p else "", p.current_year if p else "",
         u.status, u.created_at.date() if u.created_at else ""]
        for u, p, m in q.limit(MAX_EXPORT_ROWS)
    ]
    return cols, rows


def _companies(db: Session, start: date, end: date):
    q = (
        db.query(am.Company)
        .filter(func.date(am.Company.created_at) >= start)
        .filter(func.date(am.Company.created_at) <= end)
        .order_by(am.Company.name)
    )
    cols = ["Company", "Industry", "Location", "Contact", "Email",
            "Verified", "Active", "Internships", "Added On"]
    rows = [
        [c.name, c.industry or "", c.location or "", c.contact_person or "",
         c.contact_email or "", "Yes" if c.is_verified else "No",
         "Yes" if c.is_active else "No", len(c.internships),
         c.created_at.date() if c.created_at else ""]
        for c in q.limit(MAX_EXPORT_ROWS)
    ]
    return cols, rows


def _colleges(db: Session, start: date, end: date):
    q = (
        db.query(am.College)
        .filter(func.date(am.College.created_at) >= start)
        .filter(func.date(am.College.created_at) <= end)
        .order_by(am.College.name)
    )
    cols = ["College", "University", "Code", "Location", "Placement Officer",
            "Contact", "Partnership", "Added On"]
    rows = [
        [c.name, c.university or "", c.college_code or "", c.location or "",
         c.placement_officer or "", c.contact_email or "",
         c.partnership_status, c.created_at.date() if c.created_at else ""]
        for c in q.limit(MAX_EXPORT_ROWS)
    ]
    return cols, rows


def _applications(db: Session, start: date, end: date):
    q = (
        db.query(am.Application, am.Internship, am.Company, am.StudentProfile, am.User)
        .join(am.Internship, am.Internship.id == am.Application.internship_id)
        .join(am.Company, am.Company.id == am.Internship.company_id)
        .join(am.User, am.User.id == am.Application.student_id)
        .outerjoin(am.StudentProfile, am.StudentProfile.user_id == am.Application.student_id)
        .filter(func.date(am.Application.applied_at) >= start)
        .filter(func.date(am.Application.applied_at) <= end)
        .order_by(am.Application.applied_at.desc())
    )
    cols = ["Applied On", "Student", "Email", "College", "Internship",
            "Company", "Department", "Status"]
    rows = [
        [a.applied_at.date() if a.applied_at else "",
         sp.full_name if sp else "", u.email, sp.college_name if sp else "",
         i.title, c.name, i.department or "", a.status]
        for a, i, c, sp, u in q.limit(MAX_EXPORT_ROWS)
    ]
    return cols, rows


def _placements(db: Session, start: date, end: date):
    """Students who actually started -- stages 7 to 9 of the pipeline."""
    q = (
        db.query(am.Application, am.Internship, am.Company, am.StudentProfile, am.User)
        .join(am.Internship, am.Internship.id == am.Application.internship_id)
        .join(am.Company, am.Company.id == am.Internship.company_id)
        .join(am.User, am.User.id == am.Application.student_id)
        .outerjoin(am.StudentProfile, am.StudentProfile.user_id == am.Application.student_id)
        .filter(am.Application.status.in_(am.PLACED_STATUSES))
        .filter(func.date(am.Application.applied_at) >= start)
        .filter(func.date(am.Application.applied_at) <= end)
        .order_by(am.Application.updated_at.desc())
    )
    cols = ["Student", "Email", "College", "Internship", "Company",
            "Stipend", "Status"]
    rows = [
        [sp.full_name if sp else "", u.email, sp.college_name if sp else "",
         i.title, c.name, i.stipend or "", a.status]
        for a, i, c, sp, u in q.limit(MAX_EXPORT_ROWS)
    ]
    return cols, rows


def _completions(db: Session, start: date, end: date):
    q = (
        db.query(om.InternshipProgress, am.Application, am.Internship,
                 am.Company, am.StudentProfile)
        .join(am.Application, am.Application.id == om.InternshipProgress.application_id)
        .join(am.Internship, am.Internship.id == am.Application.internship_id)
        .join(am.Company, am.Company.id == am.Internship.company_id)
        .outerjoin(am.StudentProfile, am.StudentProfile.user_id == am.Application.student_id)
        .filter(om.InternshipProgress.actual_end_date.isnot(None))
        .order_by(om.InternshipProgress.actual_end_date.desc())
    )
    cols = ["Student", "College", "Internship", "Company", "Joined",
            "Ended", "Attendance %", "Completion %", "Score"]
    rows = []
    for p, a, i, c, sp in q.limit(MAX_EXPORT_ROWS):
        att = round((p.days_present or 0) / p.days_total * 100) if p.days_total else ""
        rows.append([
            sp.full_name if sp else "", sp.college_name if sp else "",
            i.title, c.name, p.joining_date or "", p.actual_end_date or "",
            att, p.completion_percent or 0,
            float(p.performance_score) if p.performance_score is not None else "",
        ])
    return cols, rows


def _college_wise(db: Session, start: date, end: date):
    rows_map: dict[str, dict[str, int]] = {}
    for name, n in (
        db.query(am.StudentProfile.college_name, func.count())
        .filter(am.StudentProfile.college_name.isnot(None))
        .group_by(am.StudentProfile.college_name).all()
    ):
        rows_map.setdefault(name or "(not given)", {})["students"] = n

    for name, n in (
        db.query(am.StudentProfile.college_name, func.count(am.Application.id))
        .join(am.Application, am.Application.student_id == am.StudentProfile.user_id)
        .group_by(am.StudentProfile.college_name).all()
    ):
        rows_map.setdefault(name or "(not given)", {})["applications"] = n

    for name, n in (
        db.query(am.StudentProfile.college_name, func.count(am.Application.id))
        .join(am.Application, am.Application.student_id == am.StudentProfile.user_id)
        .filter(am.Application.status.in_(am.PLACED_STATUSES))
        .group_by(am.StudentProfile.college_name).all()
    ):
        rows_map.setdefault(name or "(not given)", {})["placed"] = n

    cols = ["College", "Students", "Applications", "Placed", "Placement %"]
    rows = []
    for name, d in sorted(rows_map.items(), key=lambda kv: -kv[1].get("students", 0)):
        students = d.get("students", 0)
        apps = d.get("applications", 0)
        placed = d.get("placed", 0)
        pct = round(placed / students * 100, 1) if students else 0
        rows.append([name, students, apps, placed, pct])
    return cols, rows


def _company_wise(db: Session, start: date, end: date):
    cols = ["Company", "Internships", "Applications", "Selected", "Completed"]
    rows = []
    for c in db.query(am.Company).order_by(am.Company.name).all():
        ids = [i.id for i in c.internships]
        apps = selected = completed = 0
        if ids:
            apps = db.query(am.Application).filter(
                am.Application.internship_id.in_(ids)).count()
            selected = db.query(am.Application).filter(
                am.Application.internship_id.in_(ids),
                am.Application.status.in_(am.PLACED_STATUSES)).count()
            completed = db.query(am.Application).filter(
                am.Application.internship_id.in_(ids),
                am.Application.status == am.APP_COMPLETED).count()
        rows.append([c.name, len(ids), apps, selected, completed])
    return cols, rows


def _department_wise(db: Session, start: date, end: date):
    q = (
        db.query(am.Internship.department, func.count(am.Application.id))
        .outerjoin(am.Application, am.Application.internship_id == am.Internship.id)
        .group_by(am.Internship.department).all()
    )
    cols = ["Department", "Applications"]
    rows = [[d or "(not set)", n] for d, n in sorted(q, key=lambda t: -t[1])]
    return cols, rows


def _certificates(db: Session, start: date, end: date):
    q = (
        db.query(om.Certificate, am.StudentProfile, am.Internship, am.Company)
        .join(am.Application, am.Application.id == om.Certificate.application_id)
        .join(am.Internship, am.Internship.id == am.Application.internship_id)
        .join(am.Company, am.Company.id == am.Internship.company_id)
        .outerjoin(am.StudentProfile, am.StudentProfile.user_id == om.Certificate.student_id)
        .order_by(om.Certificate.created_at.desc())
    )
    cols = ["Certificate No", "Student", "Internship", "Company", "Status", "Issued On"]
    rows = [
        [c.certificate_number, sp.full_name if sp else "", i.title, co.name,
         c.status, c.issued_on or ""]
        for c, sp, i, co in q.limit(MAX_EXPORT_ROWS)
    ]
    return cols, rows


REPORTS: dict[str, tuple[str, Callable]] = {
    "students":      ("Student Registration Report", _students),
    "companies":     ("Company Registration Report", _companies),
    "colleges":      ("College Registration Report", _colleges),
    "applications":  ("Applications Report", _applications),
    "placements":    ("Placement Report", _placements),
    "completions":   ("Internship Completion Report", _completions),
    "college-wise":  ("College-wise Report", _college_wise),
    "company-wise":  ("Company-wise Report", _company_wise),
    "department":    ("Department-wise Report", _department_wise),
    "certificates":  ("Certificates Report", _certificates),
}


def _date_range(date_from: Optional[date], date_to: Optional[date]) -> tuple[date, date]:
    end = date_to or date.today()
    start = date_from or (end - timedelta(days=365))
    if start > end:
        raise HTTPException(status.HTTP_400_BAD_REQUEST,
                            "date_from cannot be after date_to")
    return start, end


@router.get("/reports", response_model=None)
def list_reports(user: am.User = Depends(perm.require_permission(om.RES_REPORTS, om.ACT_VIEW))):
    return {
        "reports": [{"key": k, "name": v[0]} for k, v in REPORTS.items()],
        "formats": ["json", "csv", "xlsx"],
        "max_rows": MAX_EXPORT_ROWS,
    }


@router.get("/reports/{report_key}", response_model=None)
def run_report(
    report_key: str,
    user: am.User = Depends(perm.require_permission(om.RES_REPORTS, om.ACT_VIEW)),
    db: Session = Depends(get_db),
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
):
    if report_key not in REPORTS:
        raise HTTPException(404, f"Unknown report. Available: {', '.join(REPORTS)}")
    name, builder = REPORTS[report_key]
    start, end = _date_range(date_from, date_to)
    columns, rows = builder(db, start, end)
    return {
        "report": report_key, "name": name,
        "date_from": start, "date_to": end,
        "columns": columns, "rows": rows, "row_count": len(rows),
        "truncated": len(rows) >= MAX_EXPORT_ROWS,
    }


@router.get("/reports/{report_key}/export")
def export_report(
    report_key: str,
    request: Request,
    fmt: str = Query("csv", pattern="^(csv|xlsx)$"),
    user: am.User = Depends(perm.require_permission(om.RES_REPORTS, om.ACT_EXPORT)),
    db: Session = Depends(get_db),
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
):
    """Download a report. Export is a separate permission from view, because
    walking out with a spreadsheet of student contact details is a different
    act from reading it on screen."""
    if report_key not in REPORTS:
        raise HTTPException(404, f"Unknown report. Available: {', '.join(REPORTS)}")
    name, builder = REPORTS[report_key]
    start, end = _date_range(date_from, date_to)
    columns, rows = builder(db, start, end)

    stamp = date.today().isoformat()
    perm.record_audit(db, user, "exported", "report", None,
                      summary=f"{name} ({fmt}, {len(rows)} rows)", request=request)

    if fmt == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(columns)
        w.writerows(rows)
        data = buf.getvalue().encode("utf-8-sig")   # BOM so Excel reads UTF-8
        return StreamingResponse(
            io.BytesIO(data), media_type="text/csv",
            headers={"Content-Disposition":
                     f'attachment; filename="{report_key}-{stamp}.csv"'},
        )

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = report_key[:31] or "Report"      # Excel caps sheet names at 31

    header_fill = PatternFill("solid", start_color="0A1F44")
    header_font = Font(color="FFFFFF", bold=True)
    ws.append(columns)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for r in rows:
        ws.append(["" if v is None else v for v in r])

    for idx, col in enumerate(columns, start=1):
        widest = max(
            [len(str(col))] + [len(str(r[idx - 1])) for r in rows[:200]] or [10]
        )
        ws.column_dimensions[get_column_letter(idx)].width = min(45, widest + 3)
    ws.freeze_panes = "A2"

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition":
                 f'attachment; filename="{report_key}-{stamp}.xlsx"'},
    )


# ===========================================================================
# 16. ANALYTICS
# ===========================================================================

@router.get("/analytics", response_model=None)
def analytics(
    user: am.User = Depends(perm.require_permission(om.RES_REPORTS, om.ACT_VIEW)),
    db: Session = Depends(get_db),
    months: int = Query(12, ge=1, le=36),
):
    """Chart-ready series. Shapes are stable so the UI can bind directly."""
    since = date.today() - timedelta(days=months * 31)

    # Postgres and SQLite spell month-truncation differently, and SQLAlchemy
    # needs the expression itself in GROUP BY -- a positional "1" is raw-SQL
    # syntax it will not accept.
    is_sqlite = db.get_bind().dialect.name == "sqlite"

    def month_of(date_col):
        return (func.strftime("%Y-%m", date_col) if is_sqlite
                else func.to_char(date_col, "YYYY-MM"))

    def by_month(date_col, extra_filter=None):
        expr = month_of(date_col)
        q = db.query(expr.label("month"), func.count()).filter(func.date(date_col) >= since)
        if extra_filter is not None:
            q = q.filter(extra_filter)
        return [
            {"month": m, "count": n}
            for m, n in q.group_by(expr).order_by(expr).all()
        ]

    student_growth = by_month(am.User.created_at, am.User.role == am.ROLE_STUDENT)
    applications = by_month(am.Application.applied_at)
    placements = by_month(
        am.Application.applied_at,
        am.Application.status.in_(am.PLACED_STATUSES),
    )

    # Skills come from a comma-separated column, so they have to be split in
    # Python. Fine at this scale; move to a join table if it gets slow.
    skill_counts: dict[str, int] = {}
    for (raw,) in db.query(am.StudentProfile.skills).filter(
        am.StudentProfile.skills.isnot(None)
    ).all():
        for s in (raw or "").split(","):
            s = s.strip()
            if s:
                skill_counts[s] = skill_counts.get(s, 0) + 1
    popular_skills = [
        {"skill": k, "count": v}
        for k, v in sorted(skill_counts.items(), key=lambda kv: -kv[1])[:15]
    ]

    status_breakdown = [
        {"status": s, "count": n}
        for s, n in db.query(am.Application.status, func.count())
        .group_by(am.Application.status).all()
    ]

    top_companies = [
        {"company": n, "applications": c}
        for n, c in db.query(am.Company.name, func.count(am.Application.id))
        .join(am.Internship, am.Internship.company_id == am.Company.id)
        .join(am.Application, am.Application.internship_id == am.Internship.id)
        .group_by(am.Company.name).order_by(func.count(am.Application.id).desc())
        .limit(10).all()
    ]

    top_colleges = [
        {"college": n, "students": c}
        for n, c in db.query(am.StudentProfile.college_name, func.count())
        .filter(am.StudentProfile.college_name.isnot(None))
        .group_by(am.StudentProfile.college_name)
        .order_by(func.count().desc()).limit(10).all()
    ]

    course_distribution = [
        {"course": c or "(not set)", "count": n}
        for c, n in db.query(am.StudentProfile.course, func.count())
        .group_by(am.StudentProfile.course).order_by(func.count().desc()).all()
    ]

    work_mode = [
        {"mode": m or "(not set)", "count": n}
        for m, n in db.query(am.Internship.work_mode, func.count())
        .group_by(am.Internship.work_mode).all()
    ]

    total_apps = sum(x["count"] for x in status_breakdown) or 0
    placed = sum(x["count"] for x in status_breakdown
                 if x["status"] in am.PLACED_STATUSES)

    return {
        "student_growth": student_growth,
        "applications_per_month": applications,
        "placements_per_month": placements,
        "application_status_breakdown": status_breakdown,
        "popular_skills": popular_skills,
        "top_companies": top_companies,
        "top_colleges": top_colleges,
        "course_distribution": course_distribution,
        "work_mode_distribution": work_mode,
        "conversion": {
            "total_applications": total_apps,
            "placed": placed,
            "placement_rate_percent": round(placed / total_apps * 100, 1) if total_apps else 0,
        },
        "note": "Gender distribution is not available -- gender is not collected "
                "anywhere in the platform.",
    }

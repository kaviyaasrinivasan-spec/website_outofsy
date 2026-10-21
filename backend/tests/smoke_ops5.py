"""Sections 13, 15, 16, 17, 20: documents, reports, analytics, approvals, settings."""
import sys, pathlib, importlib, io, csv

BACKEND = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACKEND))

db = importlib.import_module("app.database")
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

p = pathlib.Path(__file__).parent / "smoke_ops5.db"
if p.exists():
    p.unlink()
db.engine = create_engine(f"sqlite:///{p}", connect_args={"check_same_thread": False})
db.SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=db.engine)

main = importlib.import_module("app.main")
main.SessionLocal = db.SessionLocal
import app.auth_models as am
import app.ops_models as om
import app.security as sec
from app.permissions import seed_default_permissions
db.Base.metadata.create_all(bind=db.engine)

from fastapi.testclient import TestClient

passed = failed = 0
def check(label, got, want):
    global passed, failed
    ok = got == want
    if ok: passed += 1
    else:  failed += 1
    print(f"  [{'OK ' if ok else 'FAIL'}] {label:<58} {got} (want {want})")

s = db.SessionLocal()
seed_default_permissions(s)
for e, r in [("super@ops.co", om.ROLE_SUPER_ADMIN), ("mgr@ops.co", om.ROLE_OPS_MANAGER),
             ("exec@ops.co", om.ROLE_OPS_EXECUTIVE)]:
    s.add(am.User(email=e, password_hash=sec.hash_password("TestPass123"),
                  role=r, status=am.STATUS_ACTIVE, email_verified=True))
s.commit(); s.close()

def login(email):
    c = TestClient(main.app)
    assert c.post("/auth/login", json={"email": email, "password": "TestPass123"}).status_code == 200
    return c

sup, mgr, ex = login("super@ops.co"), login("mgr@ops.co"), login("exec@ops.co")

co = sup.post("/ops/companies", json={"name":"Acme Tech","industry":"Software"}).json()["id"]
col = sup.post("/ops/colleges", json={"name":"Anna University"}).json()["id"]
i1 = sup.post("/ops/internships", json={"company_id":co,"title":"Backend Intern",
                                        "department":"Engineering","work_mode":"Remote",
                                        "status":"active","openings":5}).json()["id"]
stu = TestClient(main.app)
stu.post("/auth/register", json={"full_name":"Kaviya S","email":"kaviya@oosqa.co",
                                 "password":"TestPass123","college_name":"Anna University"})
stu.patch("/student/profile", json={"skills":"Python, React, SQL","course":"B.E / B.Tech"})
sid = sup.get("/ops/students").json()["items"][0]["id"]
aid = stu.post("/student/applications", json={"internship_id":i1}).json()["id"]

print("=" * 102)
print("SECTION 13 -- DOCUMENTS")
k = sup.get("/ops/documents/kinds").json()
check("kinds listed", "resume" in k["kinds"], True)
check("  aadhaar number NOT a field", "aadhaar_number" in k["kinds"], False)
check("  note explains identity policy", "college_id" in k["note"], True)

pdf = b"%PDF-1.4\n%fake pdf for testing\n"
r = sup.post("/ops/documents", files={"file": ("resume.pdf", pdf, "application/pdf")},
             data={"owner_type":"student","owner_id":str(sid),"kind":"resume"})
check("upload resume -> 201", r.status_code, 201)
doc_id = r.json()["id"]
check("bad owner_type -> 400",
      sup.post("/ops/documents", files={"file":("a.pdf",pdf,"application/pdf")},
               data={"owner_type":"alien","owner_id":"1","kind":"resume"}).status_code, 400)
check("bad kind -> 400",
      sup.post("/ops/documents", files={"file":("a.pdf",pdf,"application/pdf")},
               data={"owner_type":"student","owner_id":str(sid),"kind":"passport"}).status_code, 400)
check("nonexistent owner -> 400",
      sup.post("/ops/documents", files={"file":("a.pdf",pdf,"application/pdf")},
               data={"owner_type":"student","owner_id":"99999","kind":"resume"}).status_code, 400)
check(".exe rejected -> 400",
      sup.post("/ops/documents", files={"file":("x.exe",b"MZ","application/octet-stream")},
               data={"owner_type":"student","owner_id":str(sid),"kind":"resume"}).status_code, 400)
check("empty file -> 400",
      sup.post("/ops/documents", files={"file":("e.pdf",b"","application/pdf")},
               data={"owner_type":"student","owner_id":str(sid),"kind":"resume"}).status_code, 400)
check("oversized -> 413",
      sup.post("/ops/documents", files={"file":("big.pdf",b"x"*(6*1024*1024),"application/pdf")},
               data={"owner_type":"student","owner_id":str(sid),"kind":"resume"}).status_code, 413)
check("list paginated", "total_pages" in sup.get("/ops/documents").json(), True)
check("  starts unverified", sup.get("/ops/documents").json()["items"][0]["is_verified"], False)
check("download -> 200", sup.get(f"/ops/documents/{doc_id}/download").status_code, 200)
check("executive CANNOT verify -> 403",
      ex.post(f"/ops/documents/{doc_id}/verify").status_code, 403)
check("manager can verify", mgr.post(f"/ops/documents/{doc_id}/verify").status_code, 200)
check("  now verified", sup.get("/ops/documents").json()["items"][0]["is_verified"], True)

print("\nSECTION 17 -- APPROVAL CENTRE")
a = sup.get("/ops/approvals").json()
check("counts present", set(a["counts"]), {"companies","colleges","internships","documents","certificates"})
check("  unverified company queued", a["counts"]["companies"], 1)
check("  pending college queued", a["counts"]["colleges"], 1)
check("  verified doc no longer queued", a["counts"]["documents"], 0)
sup.post(f"/ops/companies/{co}/verify?approve=true")
check("verifying clears the queue", sup.get("/ops/approvals").json()["counts"]["companies"], 0)

print("\nSECTION 15 -- REPORTS")
lst = sup.get("/ops/reports").json()
check("10 reports available", len(lst["reports"]), 10)
check("  formats advertised", set(lst["formats"]), {"json","csv","xlsx"})
for key in ["students","companies","colleges","applications","placements",
            "completions","college-wise","company-wise","department","certificates"]:
    r = sup.get(f"/ops/reports/{key}")
    ok = r.status_code == 200 and "columns" in r.json() and "rows" in r.json()
    check(f"  {key}", ok, True)
check("unknown report -> 404", sup.get("/ops/reports/nonsense").status_code, 404)
check("bad date range -> 400",
      sup.get("/ops/reports/students?date_from=2026-12-01&date_to=2026-01-01").status_code, 400)
rep = sup.get("/ops/reports/students").json()
check("student report has our student", rep["row_count"], 1)
check("  columns labelled", "Member ID" in rep["columns"], True)

print("\n  exports:")
# Export is granted to executives by default, but must be separately
# revocable from view -- that separation is the point of the permission.
check("executive can export by default",
      ex.get("/ops/reports/students/export?fmt=csv").status_code, 200)
check("revoke reports:export from executive",
      sup.patch("/ops/permissions", json={"role":"ops_executive","resource":"reports",
                                          "action":"export","allowed":False}).status_code, 200)
check("  export now blocked -> 403",
      ex.get("/ops/reports/students/export?fmt=csv").status_code, 403)
check("  but viewing still works", ex.get("/ops/reports/students").status_code, 200)
sup.patch("/ops/permissions", json={"role":"ops_executive","resource":"reports",
                                    "action":"export","allowed":True})
r = sup.get("/ops/reports/students/export?fmt=csv")
check("CSV export -> 200", r.status_code, 200)
check("  content-type", r.headers["content-type"].split(";")[0], "text/csv")
check("  filename set", "attachment" in r.headers["content-disposition"], True)
body = r.content.decode("utf-8-sig")
rows = list(csv.reader(io.StringIO(body)))
check("  header + 1 data row", len(rows), 2)
check("  student name present", "Kaviya S" in body, True)

r = sup.get("/ops/reports/applications/export?fmt=xlsx")
check("XLSX export -> 200", r.status_code, 200)
check("  xlsx content-type",
      "spreadsheetml" in r.headers["content-type"], True)
check("  valid zip magic (xlsx is a zip)", r.content[:2], b"PK")
from openpyxl import load_workbook
wb = load_workbook(io.BytesIO(r.content))
ws = wb.active
check("  sheet named after report", ws.title, "applications")
check("  header row written", ws.cell(1,1).value, "Applied On")
check("  frozen header", ws.freeze_panes, "A2")
check("  data row present", ws.max_row, 2)

print("\nSECTION 16 -- ANALYTICS")
an = sup.get("/ops/analytics").json()
for key in ["student_growth","applications_per_month","placements_per_month",
            "application_status_breakdown","popular_skills","top_companies",
            "top_colleges","course_distribution","work_mode_distribution","conversion"]:
    check(f"  {key} present", key in an, True)
check("skills split from CSV column", len(an["popular_skills"]), 3)
check("  Python counted", any(s["skill"]=="Python" for s in an["popular_skills"]), True)
check("conversion computed", an["conversion"]["total_applications"], 1)
check("  placement rate 0 (nobody placed yet)", an["conversion"]["placement_rate_percent"], 0)
check("gender absence stated honestly", "not collected" in an["note"], True)

print("\n  after a placement:")
for st in ["under_review","shortlisted","selected","offer_released","joined"]:
    sup.patch(f"/ops/applications/{aid}", json={"status":st})
an2 = sup.get("/ops/analytics").json()
check("placement rate now 100", an2["conversion"]["placement_rate_percent"], 100.0)
check("placements report populated", sup.get("/ops/reports/placements").json()["row_count"], 1)

print("\nSECTION 20 -- SETTINGS")
st = sup.get("/ops/settings").json()["categories"]
check("all 7 categories declared", len(st), 7)
check("  smtp_password marked secret",
      [f for f in st["email"] if f["key"]=="smtp_password"][0]["secret"], True)
check("  starts unconfigured",
      [f for f in st["email"] if f["key"]=="smtp_password"][0]["has_value"], False)
check("set a plain value",
      sup.put("/ops/settings", json={"category":"branding","key":"site_name",
                                     "value":"Out Of Syllabus"}).status_code, 200)
check("  readable back",
      [f for f in sup.get("/ops/settings").json()["categories"]["branding"]
       if f["key"]=="site_name"][0]["value"], "Out Of Syllabus")
check("unknown key -> 400",
      sup.put("/ops/settings", json={"category":"branding","key":"nope","value":"x"}).status_code, 400)
check("unknown category -> 422",
      sup.put("/ops/settings", json={"category":"aliens","key":"x","value":"y"}).status_code, 422)

print("\n  secret handling:")
check("write a secret",
      sup.put("/ops/settings", json={"category":"email","key":"smtp_password",
                                     "value":"sup3rs3cret"}).status_code, 200)
f = [f for f in sup.get("/ops/settings").json()["categories"]["email"]
     if f["key"]=="smtp_password"][0]
check("  masked on read", f["value"], "********")
check("  has_value true", f["has_value"], True)
raw = db.SessionLocal().query(om.SystemSetting).filter(
    om.SystemSetting.key=="smtp_password").first().value
check("  real value stored", raw, "sup3rs3cret")
check("writing the mask back is ignored",
      sup.put("/ops/settings", json={"category":"email","key":"smtp_password",
                                     "value":"********"}).json()["note"] is not None, True)
raw2 = db.SessionLocal().query(om.SystemSetting).filter(
    om.SystemSetting.key=="smtp_password").first().value
check("  secret NOT clobbered by the mask", raw2, "sup3rs3cret")
logs = sup.get("/ops/audit-logs?entity_type=setting").json()
check("  secret value absent from audit trail",
      all("sup3rs3cret" not in (l["summary"] or "") for l in logs["items"]), True)

print("\n  health reporting:")
h = sup.get("/ops/settings/health").json()
check("email not ready (host/user missing)", h["channels"]["email"]["configured"], False)
check("  explains what's missing", "smtp_host" in h["channels"]["email"]["note"], True)
check("whatsapp flags Meta approval", "Meta" in h["channels"]["whatsapp"]["note"], True)
check("certificate signing reported honestly",
      h["certificates"]["cryptographic_signing"], False)
for k in ["smtp_host","smtp_user"]:
    sup.put("/ops/settings", json={"category":"email","key":k,"value":"x"})
check("email ready once all three set",
      sup.get("/ops/settings/health").json()["channels"]["email"]["configured"], True)

print("\nPERMISSIONS")
check("executive cannot view settings -> 403", ex.get("/ops/settings").status_code, 403)
check("manager cannot edit settings -> 403",
      mgr.put("/ops/settings", json={"category":"branding","key":"site_name",
                                     "value":"hack"}).status_code, 403)

print("\n" + "=" * 102)
print(f"RESULT: {passed} passed, {failed} failed")
print("=" * 102)
sys.exit(1 if failed else 0)
